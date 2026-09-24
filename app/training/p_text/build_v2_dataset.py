"""Assemble reviewed P_text v2 rows offline; never train or deduplicate.

Run with ``python -m app.training.p_text.build_v2_dataset`` from the repo root.
All paths can be overridden via CLI options. Blank-content counts cover rows
that pass the source's label/approval rules; selected counts are output rows.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass, field
import json
import os
from pathlib import Path
import tempfile
from typing import Any

from openpyxl import Workbook, load_workbook


DEFAULT_BASELINE = Path("data/ReView_Integrated_Review_Dataset_v4_3_전체재검수_SUSPICIOUS확정.xlsx")
DEFAULT_EXISTING = Path("data/p_text_v2/existing_suspicious_recheck_reviewed.xlsx")
DEFAULT_MERGED = Path("data/p_text_v2/merged_ptext_labeling_workbook_final_reviewed.xlsx")
DEFAULT_OUTPUT = Path("data/p_text_v2/ptext_v2_training_master.xlsx")
OUTPUT_COLUMNS = (
    "source_set", "source_id", "content", "final_label", "use_for_training", "source_note",
)


@dataclass(frozen=True)
class Source:
    path: Path
    name: str
    sheet: str
    text_column: str
    label_column: str
    id_column: str
    requires_approval: bool


@dataclass
class SourceCounts:
    rows_read: int = 0
    selected: dict[str, int] = field(default_factory=lambda: {"NORMAL": 0, "SUSPICIOUS": 0})
    empty_content_excluded: int = 0


def _read_source(source: Source) -> tuple[list[list[Any]], SourceCounts]:
    workbook = load_workbook(source.path, read_only=True, data_only=True)
    counts = SourceCounts()
    selected = []
    try:
        if source.sheet not in workbook.sheetnames:
            raise ValueError(f"{source.path}: missing sheet {source.sheet!r}")
        rows = workbook[source.sheet].iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        required = {source.text_column, source.label_column}
        if source.requires_approval:
            required.add("use_for_training")
        missing = required.difference(headers)
        if missing:
            raise ValueError(f"{source.path}: missing required columns: {', '.join(sorted(missing))}")
        if len([name for name in headers if name]) != len({name for name in headers if name}):
            raise ValueError(f"{source.path}: duplicate column headers")
        for excel_row, values in enumerate(rows, start=2):
            counts.rows_read += 1
            row = dict(zip(headers, values, strict=True))
            label = row[source.label_column]
            allowed = ("NORMAL", "SUSPICIOUS") if source.requires_approval else ("NORMAL",)
            if label not in allowed:
                continue
            if source.requires_approval and str(row["use_for_training"]).strip().upper() != "YES":
                continue
            content = row[source.text_column]
            if content is None or not str(content).strip():
                counts.empty_content_excluded += 1
                continue
            # Only test for emptiness above: retain whitespace, case and duplicates.
            content = str(content)
            source_id = row.get(source.id_column)
            if source_id is None or not str(source_id).strip():
                source_id = f"{source.sheet}:{excel_row}"
            selected.append([
                source.name, source_id, content, label, "YES",
                f"{source.path.name}; sheet={source.sheet}; row={excel_row}",
            ])
            counts.selected[label] += 1
    finally:
        workbook.close()
    return selected, counts


def _publish(rows: list[list[Any]], output: Path) -> None:
    """Publish a completed same-directory file atomically, without replacement."""
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "labeling"
    temporary: Path | None = None
    try:
        for row_number, values in enumerate([OUTPUT_COLUMNS, *rows], start=1):
            for column, value in enumerate(values, start=1):
                if isinstance(value, str) and len(value) > 32767:
                    raise ValueError(f"Output row {row_number}: text exceeds Excel cell limit")
                cell = sheet.cell(row=row_number, column=column, value=value)
                # Preserve literal review text such as '=great', not Excel formulas.
                if isinstance(value, str):
                    cell.data_type = "s"
        with tempfile.NamedTemporaryFile(
            dir=output.parent, prefix=f".{output.stem}-", suffix=".xlsx", delete=False,
        ) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        # Unlike replace/rename on some platforms, link fails if output exists,
        # including when another process created it after the initial check.
        os.link(temporary, output)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def build_v2_dataset(
    baseline_path: str | Path = DEFAULT_BASELINE,
    existing_path: str | Path = DEFAULT_EXISTING,
    merged_path: str | Path = DEFAULT_MERGED,
    output_path: str | Path = DEFAULT_OUTPUT,
) -> dict[str, Any]:
    """Build a new workbook and return audit counts. Existing outputs are refused."""
    output = Path(output_path)
    sources = [
        Source(Path(baseline_path), "baseline_normal", "REVIEW_MASTER", "review", "final_label", "master_id", False),
        Source(Path(existing_path), "existing_rechecked", "recheck", "review", "new_final_label", "master_id", True),
        Source(Path(merged_path), "merged_rechecked", "labeling", "content", "final_label", "source_excel_row", True),
    ]
    if os.path.lexists(output):
        raise FileExistsError(f"Refusing to overwrite existing output: {output}")
    if any(output.resolve() == source.path.resolve() for source in sources):
        raise ValueError("Output path must differ from every input path")
    rows = []
    source_counts = {}
    totals = {"NORMAL": 0, "SUSPICIOUS": 0}
    empty_count = 0
    for source in sources:
        selected, counts = _read_source(source)
        rows.extend(selected)
        source_counts[source.name] = {
            "rows_read": counts.rows_read,
            "selected": counts.selected,
            "empty_content_excluded": counts.empty_content_excluded,
        }
        empty_count += counts.empty_content_excluded
        for label, count in counts.selected.items():
            totals[label] += count
    _publish(rows, output)
    return {
        "sources": source_counts,
        "empty_content_excluded": empty_count,
        "output_rows": len(rows),
        "NORMAL": totals["NORMAL"],
        "SUSPICIOUS": totals["SUSPICIOUS"],
    }


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline", type=Path, default=DEFAULT_BASELINE)
    parser.add_argument("--existing", type=Path, default=DEFAULT_EXISTING)
    parser.add_argument("--merged", type=Path, default=DEFAULT_MERGED)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    counts = build_v2_dataset(args.baseline, args.existing, args.merged, args.output)
    print(json.dumps(counts, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
