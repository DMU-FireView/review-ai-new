import json
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

from app.training.p_text import build_v2_dataset as builder
from app.training.p_text.data import prepare_records, read_ptext_v2_workbook


def write_book(path, sheet, headers, rows):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.append(headers)
    for values in rows:
        worksheet.append(values)
        for cell in worksheet[worksheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    workbook.save(path)
    workbook.close()


@pytest.fixture
def inputs(tmp_path):
    baseline = tmp_path / "baseline.xlsx"
    existing = tmp_path / "existing.xlsx"
    merged = tmp_path / "merged.xlsx"
    write_book(baseline, "REVIEW_MASTER", ["master_id", "review", "final_label"], [
        ["b1", " Same\nTEXT ", "NORMAL"],
        ["b2", " Same\nTEXT ", "NORMAL"],
        ["b3", "baseline suspicious excluded", "SUSPICIOUS"],
        ["b4", "uncertain", "UNCERTAIN"],
        ["b5", None, "NORMAL"],
        ["b6", " \t ", "NORMAL"],
    ])
    for path, sheet, text_column, label_column, id_column in [
        (existing, "recheck", "review", "new_final_label", "master_id"),
        (merged, "labeling", "content", "final_label", "source_excel_row"),
    ]:
        write_book(path, sheet, [id_column, text_column, label_column, "use_for_training", "reviewer_note"], [
            [101, " Same\nTEXT ", "SUSPICIOUS", "YES", "secret metadata"],
            [102, "=literal review", "NORMAL", " yes ", "secret metadata"],
            [103, "rejected", "NORMAL", "NO", "note"],
            [104, "uncertain", "UNCERTAIN", "YES", "note"],
            [105, "invalid", "INVALID", "YES", "note"],
            [106, "promotion", "DISCLOSED_PROMO", "YES", "note"],
            [107, None, "NORMAL", "YES", "note"],
            [108, " \n ", "SUSPICIOUS", "YES", "note"],
            [109, "missing approval", "NORMAL", None, "note"],
        ])
    return baseline, existing, merged


def test_assembly_rules_counts_and_source_immutability(inputs, tmp_path):
    original = [path.read_bytes() for path in inputs]
    output = tmp_path / "result.xlsx"
    with patch.object(builder, "load_workbook", wraps=openpyxl.load_workbook) as loader:
        counts = builder.build_v2_dataset(*inputs, output)
    assert loader.call_count == 3
    for call, path in zip(loader.call_args_list, inputs):
        assert call.args == (path,)
        assert call.kwargs == {"read_only": True, "data_only": True}
    assert [path.read_bytes() for path in inputs] == original
    rows = read_ptext_v2_workbook(str(output))
    assert len(rows) == 6
    assert all(tuple(row) == builder.OUTPUT_COLUMNS for row in rows)
    assert [row["source_set"] for row in rows] == [
        "baseline_normal", "baseline_normal", "existing_rechecked",
        "existing_rechecked", "merged_rechecked", "merged_rechecked",
    ]
    assert [row["source_id"] for row in rows] == ["b1", "b2", 101, 102, 101, 102]
    assert [row["final_label"] for row in rows] == [
        "NORMAL", "NORMAL", "SUSPICIOUS", "NORMAL", "SUSPICIOUS", "NORMAL",
    ]
    assert [row["content"] for row in rows] == [
        " Same\nTEXT ", " Same\nTEXT ", " Same\nTEXT ",
        "=literal review", " Same\nTEXT ", "=literal review",
    ]
    assert all(row["use_for_training"] == "YES" for row in rows)
    assert "baseline.xlsx; sheet=REVIEW_MASTER; row=2" == rows[0]["source_note"]
    assert counts == {
        "sources": {
            "baseline_normal": {"rows_read": 6, "selected": {"NORMAL": 2, "SUSPICIOUS": 0}, "empty_content_excluded": 2},
            "existing_rechecked": {"rows_read": 9, "selected": {"NORMAL": 1, "SUSPICIOUS": 1}, "empty_content_excluded": 2},
            "merged_rechecked": {"rows_read": 9, "selected": {"NORMAL": 1, "SUSPICIOUS": 1}, "empty_content_excluded": 2},
        },
        "empty_content_excluded": 6, "output_rows": 6, "NORMAL": 4, "SUSPICIOUS": 2,
    }
    prepared = prepare_records(rows, text_column="content", use_for_training_column="use_for_training")
    assert prepared.examples == [{"text": "=literal review", "label": 0}]
    assert prepared.conflicting_groups_removed == 1
    assert not list(tmp_path.glob(".result-*.xlsx"))


def test_existing_output_stops_before_reading_inputs(inputs, tmp_path):
    output = tmp_path / "result.xlsx"
    output.write_bytes(b"existing output")
    with patch.object(builder, "load_workbook") as loader:
        with pytest.raises(FileExistsError, match="Refusing to overwrite"):
            builder.build_v2_dataset(*inputs, output)
    loader.assert_not_called()
    assert output.read_bytes() == b"existing output"


def test_output_cannot_overwrite_source(inputs):
    original = inputs[0].read_bytes()
    with pytest.raises(FileExistsError):
        builder.build_v2_dataset(*inputs, inputs[0])
    assert inputs[0].read_bytes() == original


def test_publish_race_does_not_overwrite(inputs, tmp_path):
    output = tmp_path / "result.xlsx"
    real_link = builder.os.link

    def competing_publish(source, destination):
        # The temporary file is already a complete, readable workbook.
        assert len(read_ptext_v2_workbook(str(source))) == 6
        Path(destination).write_bytes(b"another writer")
        real_link(source, destination)

    with patch.object(builder.os, "link", side_effect=competing_publish):
        with pytest.raises(FileExistsError):
            builder.build_v2_dataset(*inputs, output)
    assert output.read_bytes() == b"another writer"
    assert not list(tmp_path.glob(".result-*.xlsx"))


def test_failed_save_is_not_published(inputs, tmp_path):
    output = tmp_path / "result.xlsx"
    with patch.object(builder.Workbook, "save", side_effect=OSError("save failed")):
        with pytest.raises(OSError, match="save failed"):
            builder.build_v2_dataset(*inputs, output)
    assert not output.exists()
    assert not list(tmp_path.glob(".result-*.xlsx"))


@pytest.mark.parametrize("sheet,headers", [
    ("wrong", ["review", "final_label"]),
    ("REVIEW_MASTER", ["review"]),
    ("REVIEW_MASTER", ["review", "final_label", "final_label"]),
])
def test_invalid_schema_does_not_publish(inputs, tmp_path, sheet, headers):
    write_book(inputs[0], sheet, headers, [])
    output = tmp_path / "result.xlsx"
    with pytest.raises(ValueError, match="missing|duplicate"):
        builder.build_v2_dataset(*inputs, output)
    assert not output.exists()


def test_missing_identifier_has_auditable_fallback(inputs, tmp_path):
    write_book(inputs[0], "REVIEW_MASTER", ["review", "final_label"], [["text", "NORMAL"]])
    output = tmp_path / "result.xlsx"
    builder.build_v2_dataset(*inputs, output)
    assert read_ptext_v2_workbook(str(output))[0]["source_id"] == "REVIEW_MASTER:2"


def test_cli_prints_counts_with_synthetic_inputs(inputs, tmp_path, capsys):
    builder.main([
        "--baseline", str(inputs[0]), "--existing", str(inputs[1]),
        "--merged", str(inputs[2]), "--output", str(tmp_path / "result.xlsx"),
    ])
    counts = json.loads(capsys.readouterr().out)
    assert counts["output_rows"] == 6
    assert counts["NORMAL"] == 4
    assert counts["SUSPICIOUS"] == 2
