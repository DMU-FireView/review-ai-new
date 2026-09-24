from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

from app.training.p_text.data import (
    InsufficientDataError,
    prepare_records,
    read_ptext_v2_workbook,
    read_review_master,
    stratified_split,
)


def make_workbook(tmp_path: Path, sheet: str, headers: list, rows: list) -> Path:
    path = tmp_path / "synthetic.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    if headers:
        worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    workbook.close()
    return path


def test_v2_reads_labeling_without_modifying_workbook(tmp_path: Path) -> None:
    headers = ["content", "final_label", "use_for_training", "reviewer_note"]
    values = ["review text", "NORMAL", "YES", "human metadata"]
    path = make_workbook(tmp_path, "labeling", headers, [values])
    original = path.read_bytes()
    with patch("openpyxl.load_workbook", wraps=openpyxl.load_workbook) as loader:
        assert read_ptext_v2_workbook(str(path)) == [dict(zip(headers, values))]
    loader.assert_called_once_with(str(path), read_only=True, data_only=True)
    assert path.read_bytes() == original


def test_v2_requires_labeling_sheet(tmp_path: Path) -> None:
    path = make_workbook(tmp_path, "REVIEW_MASTER", ["review"], [])
    with pytest.raises(ValueError, match="must contain a labeling sheet"):
        read_ptext_v2_workbook(str(path))


@pytest.mark.parametrize("missing", ["content", "final_label", "use_for_training"])
def test_v2_requires_each_column(tmp_path: Path, missing: str) -> None:
    headers = [name for name in ("content", "final_label", "use_for_training") if name != missing]
    path = make_workbook(tmp_path, "labeling", headers, [])
    with pytest.raises(ValueError, match=f"missing required columns: {missing}"):
        read_ptext_v2_workbook(str(path))


def test_v2_empty_sheet_has_clear_error(tmp_path: Path) -> None:
    path = make_workbook(tmp_path, "labeling", [], [])
    with pytest.raises(ValueError, match="missing required columns"):
        read_ptext_v2_workbook(str(path))


def test_v2_headers_without_rows(tmp_path: Path) -> None:
    path = make_workbook(tmp_path, "labeling", ["content", "final_label", "use_for_training"], [])
    assert read_ptext_v2_workbook(str(path)) == []


def test_v2_filtering_and_metadata_exclusion(tmp_path: Path) -> None:
    metadata = {
        "source_excel_row": 2, "platform": "test", "product_id": "p1",
        "review_id": "r1", "rating": 5, "candidate_reasons": "reason",
        "matched_features": "feature", "confirmed_patterns": "pattern",
        "reviewer_note": "note",
    }
    rows = [
        ["normal", "NORMAL", "YES"],
        ["suspicious", " suspicious ", " yes "],
        ["mixed case", "NORMAL", "YeS"],
        ["no", "NORMAL", "NO"],
        ["uncertain", "UNCERTAIN", "YES"],
        ["invalid", "INVALID", "YES"],
        ["promo", "DISCLOSED_PROMO", "YES"],
        ["unknown", "UNKNOWN", "YES"],
        ["blank flag", "NORMAL", None],
        ["boolean flag", "NORMAL", True],
        ["numeric flag", "NORMAL", 1],
        ["   ", "NORMAL", "YES"],
    ]
    path = make_workbook(
        tmp_path, "labeling", ["content", "final_label", "use_for_training", *metadata],
        [row + list(metadata.values()) for row in rows],
    )
    prepared = prepare_records(
        read_ptext_v2_workbook(str(path)), text_column="content",
        label_column="final_label", use_for_training_column="use_for_training",
    )
    assert prepared.examples == [
        {"text": "normal", "label": 0}, {"text": "suspicious", "label": 1},
        {"text": "mixed case", "label": 0},
    ]
    assert prepared.raw_selected_count == 3
    assert all(set(example) == {"text", "label"} for example in prepared.examples)


def test_v2_preserves_duplicate_and_conflict_policy() -> None:
    rows = [
        {"content": text, "final_label": label, "use_for_training": flag}
        for text, label, flag in [
            (" Same\ntext ", "NORMAL", "YES"),
            ("same TEXT", "NORMAL", "YES"),
            ("same text", "SUSPICIOUS", "NO"),
            ("CONFLICT text", "NORMAL", "YES"),
            (" conflict\ttext ", "SUSPICIOUS", "YES"),
            ("conflict text", "NORMAL", "YES"),
            ("unique", "SUSPICIOUS", "YES"),
            ("unique", "INVALID", "YES"),
        ]
    ]
    prepared = prepare_records(rows, text_column="content", use_for_training_column="use_for_training")
    assert prepared.examples == [{"text": "Same\ntext", "label": 0}, {"text": "unique", "label": 1}]
    assert prepared.raw_selected_count == 6
    assert prepared.duplicate_rows_removed == 4
    assert prepared.conflicting_groups_removed == 1
    with pytest.raises(InsufficientDataError, match="at least 3 unique examples"):
        stratified_split(prepared.examples)


def test_baseline_loader_and_default_preparation_are_unchanged(tmp_path: Path) -> None:
    headers = ["review", "final_label", "use_for_training", "platform"]
    values = [
        ["normal", "NORMAL", "NO", "test"],
        ["suspicious", "SUSPICIOUS", None, "test"],
        [" NORMAL ", "NORMAL", "YES", "test"],
        ["conflict", "NORMAL", "YES", "test"],
        ["CONFLICT", "SUSPICIOUS", "NO", "test"],
        ["skip", "UNCERTAIN", "YES", "test"],
    ]
    path = make_workbook(tmp_path, "REVIEW_MASTER", headers, values)
    records = read_review_master(str(path))
    assert records == [dict(zip(headers, row)) for row in values]
    prepared = prepare_records(records)
    assert prepared == prepare_records(records, use_for_training_column=None)
    assert prepared.examples == [{"text": "normal", "label": 0}, {"text": "suspicious", "label": 1}]
    assert prepared.raw_selected_count == 5
    assert prepared.duplicate_rows_removed == 3
    assert prepared.conflicting_groups_removed == 1


def test_optional_filter_supports_custom_column_and_value() -> None:
    rows = [
        {"review": "included", "final_label": "NORMAL", "approved": " include "},
        {"review": "excluded", "final_label": "NORMAL", "approved": "YES"},
        {"review": "missing flag", "final_label": "NORMAL"},
    ]
    prepared = prepare_records(rows, use_for_training_column="approved", use_for_training_value=" Include ")
    assert prepared.examples == [{"text": "included", "label": 0}]
